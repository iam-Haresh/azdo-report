"""End-to-end: sample raw cache -> analysis -> workbook, dossiers, summaries, diff (no network)."""

from __future__ import annotations

import openpyxl

from azdo_analyzer.analyze import run_analysis
from azdo_analyzer.report import write_reports
from tests.fixtures.sample_data import NOW, PROJECT

EXPECTED_SHEETS = [
    "Summary", "Tracker", "Repos", "Pipelines", "Releases", "Release Stages", "Tasks", "Variables", "Variable Groups",
    "Service Connections", "Secure Files", "Agents & Pools", "Environments & Deploy Groups", "Branch Policies",
    "Dependencies", "Findings", "Archetypes", "Coverage",
]


def _run(cfg, rules, provider):
    result = run_analysis(cfg, rules, [PROJECT], provider=provider, now=NOW)
    return result, write_reports(result, "contoso-retail", dossiers=True)


def test_analysis_results(sample_cfg, rules, provider):
    result, _ = _run(sample_cfg, rules, provider)
    project = result.projects[0]
    repos = {r.name: r for r in project.repos}
    pipelines = {r.pipeline.name: r for r in project.pipelines.values()}

    # YAML: cross-repo template resolved at the pinned tag (Maven@3, not the default branch's Maven@4)
    pay = pipelines["payments-api-ci"]
    assert pay.facts.tasks == ["Docker@2", "HelmDeploy@0", "Maven@3", "PublishTestResults@2", "SonarQubePrepare@5"]
    assert pay.pipeline.flags["timeline_only_tasks"] == [] and pay.pipeline.flags["static_only_tasks"] == []
    assert pay.facts.approvals and pay.facts.gates  # checks on environment payments-prod (resource-scoped name)
    assert "aks-internal" in pay.facts.service_connections and pay.facts.internal_targets

    # repo <-> pipeline links (the checklist's "link it to the project which is using these")
    assert repos["legacy-web"].pipelines == ["build:Contoso Retail:12"]
    assert repos["legacy-web"].releases == ["release:Contoso Retail:5"]
    assert repos["pipeline-templates"].template_for == ["payments-api-ci"]

    assert repos["payments-api"].profile.key == "java-maven | maven | kubernetes | image"
    assert repos["legacy-web"].profile.name == ".NET Framework -> IIS (Windows)"
    assert repos["legacy-web"].complexity == "High" and repos["old-tools"].status == "Dead"
    assert repos["Data Scripts"].status == "Disabled" and repos["Data Scripts"].findings == []

    ids = {(f.id, f.repo, f.pipeline) for f in project.all_findings()}
    expected = {
        ("DEPLOYMENT_GROUP", "legacy-web", "legacy-web-release"),
        ("KEYVAULT", "legacy-web", "legacy-web-release"),
        ("WINDOWS_TOOLCHAIN", "legacy-web", "legacy-web-CI"),
        ("TASK_GROUP", "legacy-web", "legacy-web-CI"),
        ("UNMAPPED_TASK", "legacy-web", "legacy-web-CI"),
        ("PLAINTEXT_SECRET", "legacy-web", "legacy-web-CI"),
        ("TEST_FORMAT", "legacy-web", "legacy-web-CI"),
        ("CONFIG_TRANSFORM", "legacy-web", "legacy-web-release"),
        ("HIDDEN_UI_CONFIG", "payments-api", "payments-api-ci"),
        ("CROSS_PIPELINE_TRIGGER", "payments-api", "payments-api-ci"),
        ("FEED_DEPENDENCY", "payments-api", None),
        ("FEED_DEPENDENCY", "legacy-web", None),
        ("VSO_LOGGING_CMD", "payments-api", None),
        ("ULTIMATE_ONLY", "payments-api", None),
        ("SUBMODULES", "legacy-web", None),
        ("REPO_SIZE", "legacy-web", None),
        ("ORPHAN", None, "orphan-ci"),
        ("UNUSED_CONNECTION", None, None),
    }
    missing = {e for e in expected if e not in ids}
    assert not missing, missing
    unused = [f.detail for f in project.all_findings() if f.id == "UNUSED_CONNECTION"]
    assert unused == ["old-ftp"]  # azure-prod counts as used through the Key Vault group
    assert {g["name"]: g["proposed_level"] for g in project.groups} == {"payments-common": "parent GitLab group", "legacy-kv": "external secret store"}


def test_reports_written(sample_cfg, rules, provider):
    result, paths = _run(sample_cfg, rules, provider)
    wb = openpyxl.load_workbook(paths["workbook"])
    assert wb.sheetnames == EXPECTED_SHEETS
    tracker = wb["Tracker"]
    header = [c.value for c in tracker[1]]
    rows = {r[header.index("Repo")]: dict(zip(header, r)) for r in tracker.iter_rows(min_row=2, values_only=True)}
    assert rows["legacy-web"]["Complexity"] == "High" and rows["legacy-web"]["Pipeline type"] == "Classic build, Classic release"
    assert rows["pipeline-templates"]["Pipeline type"] == "templates for payments-api-ci"
    dossier = paths["dossiers"]["Contoso Retail/legacy-web"].read_text()
    for needle in ("## 1. Project level", "## 10. Cutover and validation", "IISWebAppDeploymentOnMachineGroup@0", "web-prod-servers",
                   "Alice Admin, Bob Boss", "BUILD_PLATFORM", "```mermaid"):
        assert needle in dossier, needle
    assert "_manual_- [x]" not in dossier  # bullets must not be glued together
    summary = paths["summaries"][0].read_text()
    assert "Wave readiness" in summary and "old-tools" in summary
    assert "First run" in paths["diff"].read_text()


def test_second_run_diff_and_overrides(sample_cfg, rules, provider):
    _run(sample_cfg, rules, provider)
    sample_cfg.overrides_path.write_text(
        "repos:\n  Contoso Retail/legacy-web:\n    wave: 2\n    owner: team-web\n    status: In analysis\n    archetype: Legacy IIS\n",
        encoding="utf-8")
    result, paths = _run(sample_cfg, rules, provider)
    assert "_No changes._" in paths["diff"].read_text()
    legacy = next(r for r in result.projects[0].repos if r.name == "legacy-web")
    assert legacy.archetype == "Legacy IIS"
    tracker = openpyxl.load_workbook(paths["workbook"])["Tracker"]
    header = [c.value for c in tracker[1]]
    row = next(dict(zip(header, r)) for r in tracker.iter_rows(min_row=2, values_only=True) if r[header.index("Repo")] == "legacy-web")
    assert (row["Wave"], row["Owner"], row["Status"], row["Archetype"]) == ("2", "team-web", "In analysis", "Legacy IIS")


def test_inventory_mode_without_full_definitions(sample_cfg, rules, provider):
    raw = sample_cfg.raw_dir / PROJECT
    for folder in ("build_definition", "release_definition"):
        for f in (raw / folder).glob("*.json"):
            f.unlink()
    meta = raw / "_meta.json"
    meta.write_text('{"fetched_at": "2026-10-01T08:00:00Z", "depth": "inventory", "project_id": "6ce954b1-ce1f-45d1-b94d-e6bf2464ba2c"}')
    result = run_analysis(sample_cfg, rules, [PROJECT], provider=provider, now=NOW)
    project = result.projects[0]
    assert not project.deep
    yaml = next(r for r in project.pipelines.values() if r.pipeline.kind == "yaml")
    assert yaml.pipeline.depth == "inventory" and yaml.pipeline.flags["ui_variables"] == 2
    assert any(f.id == "INVENTORY_ONLY" for f in project.findings)
    paths = write_reports(result, "inventory", dossiers=False)
    assert paths["dossiers"] == {} and paths["workbook"].exists()
