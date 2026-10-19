"""The report docs explain every workbook column and every finding, and their links work."""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils.cell import range_boundaries

from azdo_analyzer.analyze import run_analysis
from azdo_analyzer.config import PROJECT_ROOT
from azdo_analyzer.report import write_reports
from tests.fixtures.sample_data import NOW, PROJECT

DOCS = PROJECT_ROOT / "docs" / "reports"


def sheet_doc(sheet: str) -> Path:
    return DOCS / "workbook" / (sheet.lower().replace(" & ", "-and-").replace(" ", "-") + ".md")


def test_every_workbook_column_is_documented(sample_cfg, rules, provider):
    result = run_analysis(sample_cfg, rules, [PROJECT], provider=provider, now=NOW)
    workbook = openpyxl.load_workbook(write_reports(result, "docs", dossiers=False)["workbook"])
    index = (DOCS / "workbook" / "README.md").read_text(encoding="utf-8")
    missing = []
    for ws in workbook.worksheets:
        doc = sheet_doc(ws.title)
        assert doc.exists(), f"no doc for sheet {ws.title}: expected {doc.name}"
        assert f"| {ws.title} |" in index, f"sheet {ws.title} missing from workbook/README.md"
        text = doc.read_text(encoding="utf-8")
        for table in ws.tables.values():
            min_col, min_row, max_col, _ = range_boundaries(table.ref)
            for col in range(min_col, max_col + 1):
                header = str(ws.cell(row=min_row, column=col).value)
                cell = header.replace("|", "\\|")
                if not re.search(rf"^\| {re.escape(cell)} \|", text, re.M):
                    missing.append(f"{ws.title} → {header}")
    assert missing == []


def test_findings_catalogue_matches_rules():
    catalogue = yaml.safe_load((PROJECT_ROOT / "rules" / "findings.yaml").read_text(encoding="utf-8"))["findings"]
    text = (DOCS / "workbook" / "findings.md").read_text(encoding="utf-8")
    documented = set(re.findall(r"^\| `([A-Z_]+)` \|", text, re.M))
    assert documented == set(catalogue), (set(catalogue) - documented, documented - set(catalogue))
    for fid, meta in catalogue.items():
        assert re.search(rf"^\| `{fid}` \| {meta['severity']}\b", text, re.M), f"{fid}: severity should be {meta['severity']}"


def _slug(heading: str) -> str:
    """GitHub-style heading anchor."""
    text = re.sub(r"[^\w\- ]", "", heading.strip().lower())
    return text.replace(" ", "-")


def test_relative_links_and_anchors_resolve():
    broken = []
    for doc in sorted(DOCS.rglob("*.md")):
        text = re.sub(r"`[^`\n]*`", "", doc.read_text(encoding="utf-8"))  # code spans are not links (regexes contain "](")
        for target in re.findall(r"\]\(([^)\s]+)\)", text):
            if re.match(r"^[a-z]+://", target):
                continue
            path_part, _, anchor = target.partition("#")
            resolved = (doc.parent / path_part).resolve() if path_part else doc
            if not resolved.exists():
                broken.append(f"{doc.relative_to(DOCS)} → {target}")
                continue
            if anchor and resolved.suffix == ".md":
                headings = {_slug(h) for h in re.findall(r"^#{1,6} (.+)$", resolved.read_text(encoding="utf-8"), re.M)}
                if anchor not in headings:
                    broken.append(f"{doc.relative_to(DOCS)} → {target} (no such heading)")
    assert broken == []


def test_archetype_values_page_matches_the_rules():
    from azdo_analyzer.report.archetype_doc import DOC_PATH, RULES_PATH, render

    expected = render(yaml.safe_load(RULES_PATH.read_text(encoding="utf-8")))
    assert DOC_PATH.read_text(encoding="utf-8") == expected, "regenerate it: python -m azdo_analyzer.report.archetype_doc"
